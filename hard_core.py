"""Local HARD workspace. No network calls; review copies are separate from explicit file-management actions."""
from __future__ import annotations
from contextlib import contextmanager
import base64
import hashlib
import io
import json
import os
from pathlib import Path
import re
import shutil
import sqlite3
import subprocess
import threading
import time
import uuid
import zipfile
from xml.etree import ElementTree as ET
from docx import Document
from pypdf import PdfReader

MAX_FILE = 20 * 1024 * 1024
EXTENSIONS = {'.docx', '.pdf', '.txt'}
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

class UserError(Exception):
    pass

def digest(data):
    return hashlib.sha256(data).hexdigest()

def identifier():
    return uuid.uuid4().hex

def text(value, name, maximum=500, required=True):
    if not isinstance(value, str) or len(value) > maximum or (required and not value.strip()):
        raise UserError(f'Enter a valid {name} (up to {maximum} characters).')
    return value.strip()

def safe_name(value):
    value = text(value, 'file name', 180)
    if value in {'.', '..'} or re.search(r'[\\/:*?"<>|\x00-\x1f]', value) or value.endswith(('.', ' ')):
        raise UserError('Use a file name without slashes or special characters.')
    if value.split('.')[0].upper() in {'CON','PRN','AUX','NUL',*(f'COM{i}' for i in range(1,10)),*(f'LPT{i}' for i in range(1,10))}:
        raise UserError('That name is reserved by Windows.')
    return value

def validate_docx(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 3000 or sum(e.file_size for e in entries) > 80 * 1024 * 1024:
                raise UserError('This Word document is too large when expanded.')
            if any(e.file_size > 15 * 1024 * 1024 for e in entries):
                raise UserError('This document contains an oversized component.')
            if 'word/document.xml' not in archive.namelist():
                raise UserError('This is not a valid Word document.')
            for entry in entries:
                if 'vbaproject' in entry.filename.lower() or '/embeddings/' in entry.filename.lower():
                    raise UserError('Documents with macros or embedded objects must be reviewed in Word first.')
                if entry.filename.endswith('.rels'):
                    root = ET.fromstring(archive.read(entry))
                    for rel in root:
                        if rel.get('TargetMode') == 'External' and not rel.get('Type','').endswith('/hyperlink'):
                            raise UserError('This document contains external linked content. Remove it in Word before importing.')
    except (zipfile.BadZipFile, ET.ParseError, KeyError) as exc:
        raise UserError('The Word document is damaged or unsupported.') from exc

def extract(data, suffix):
    if suffix == '.docx':
        validate_docx(data)
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            paragraphs = []
            tracked = comments = False
            for name in archive.namelist():
                if re.fullmatch(r'word/(document|header\d+|footer\d+|footnotes|endnotes)\.xml', name):
                    root = ET.fromstring(archive.read(name))
                    tracked |= next(root.iter(W+'ins'), None) is not None or next(root.iter(W+'del'), None) is not None
                    for p in root.iter(W+'p'):
                        value = ''.join(n.text or '' for n in p.iter(W+'t'))
                        if value.strip(): paragraphs.append(value)
                if name == 'word/comments.xml':
                    comments = bool(list(ET.fromstring(archive.read(name))))
        return '\n'.join(paragraphs), {'tracked_changes':tracked,'comments':comments}
    if suffix == '.pdf':
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted: raise UserError('Unlock this PDF before importing it.')
            if len(reader.pages)>250: raise UserError('Please review documents of 250 pages or fewer.')
            return '\n'.join(p.extract_text() or '' for p in reader.pages), {'pages':len(reader.pages)}
        except UserError: raise
        except Exception as exc: raise UserError('The PDF could not be read.') from exc
    try: return data.decode('utf-8-sig'), {}
    except UnicodeDecodeError as exc: raise UserError('Save this text file using UTF-8 encoding.') from exc

def findings(content, meta):
    result = []
    def add(kind, title, detail, severity='check'):
        result.append({'id':digest((kind+detail).encode())[:20], 'kind':kind,'title':title,'detail':detail,'severity':severity})
    for match in re.finditer(r'\[(?:INSERT|TODO|TBD|CLIENT|DATE|NAME|AMOUNT)[^\]\n]{0,100}\]|\b(?:TODO|TBD|XXX)\b', content, re.I):
        add('placeholder','Unfinished placeholder',match.group())
    for match in re.finditer(r'\b([A-Za-z]{2,})\s+\1\b',content,re.I):
        add('repeated_word','Possible repeated word',match.group())
    if meta.get('tracked_changes'): add('tracked_changes','Tracked changes remain','Resolve tracked changes in Word, then import the revised file.','block')
    if meta.get('comments'): add('comments','Comments remain','Check whether reviewer comments should be removed before sharing.','block')
    emails = sorted(set(re.findall(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}',content)))
    if emails: add('addresses','Check addresses in the document',', '.join(emails))
    currencies = sorted(set(re.findall(r'\b(?:USD|EUR|GBP|GHS|KES|UGX|ZAR)\b',content)))
    if len(currencies)>1: add('currency','More than one currency',', '.join(currencies)+'. Confirm that this is intentional.')
    if not content.strip(): add('no_text','No readable text','This may be a scanned PDF. OCR is not connected; review it visually.','block')
    return list({x['id']:x for x in result}.values())[:200]

class Workspace:
    def __init__(self, directory):
        self.directory = Path(directory).resolve()
        self.directory.mkdir(parents=True, exist_ok=True)
        self.blobs = self.directory / 'documents'
        self.blobs.mkdir(exist_ok=True)
        self.database = self.directory / 'workspace.sqlite3'
        self.conversion_lock = threading.Lock()
        with self.db() as db:
            db.executescript('''
            PRAGMA journal_mode=WAL;
            CREATE TABLE IF NOT EXISTS documents(id TEXT PRIMARY KEY,name TEXT NOT NULL,suffix TEXT NOT NULL,hash TEXT NOT NULL UNIQUE,content TEXT NOT NULL,meta TEXT NOT NULL,issues TEXT NOT NULL,reviewed INTEGER DEFAULT 0,pdf_id TEXT,created REAL NOT NULL);
            CREATE TABLE IF NOT EXISTS decisions(document_id TEXT,issue_id TEXT,note TEXT,PRIMARY KEY(document_id,issue_id));
            CREATE TABLE IF NOT EXISTS folders(id TEXT PRIMARY KEY,path TEXT NOT NULL UNIQUE,label TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS files(id TEXT PRIMARY KEY,folder_id TEXT,path TEXT NOT NULL UNIQUE,name TEXT NOT NULL,modified REAL NOT NULL,size INTEGER NOT NULL);
            CREATE TABLE IF NOT EXISTS contacts(id TEXT PRIMARY KEY,name TEXT NOT NULL,email TEXT NOT NULL UNIQUE,organization TEXT NOT NULL,source TEXT NOT NULL,verified REAL NOT NULL);
            CREATE TABLE IF NOT EXISTS drafts(id TEXT PRIMARY KEY,payload TEXT NOT NULL,status TEXT NOT NULL DEFAULT 'draft',created REAL NOT NULL,provider_id TEXT);
            CREATE TABLE IF NOT EXISTS events(id TEXT PRIMARY KEY,kind TEXT NOT NULL,detail TEXT NOT NULL,created REAL NOT NULL);
            CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT NOT NULL);
            ''')
            db.execute("UPDATE drafts SET status='uncertain' WHERE status='sending'")

    @contextmanager
    def db(self):
        connection = sqlite3.connect(self.database, timeout=20)
        connection.row_factory = sqlite3.Row
        try:
            with connection:
                yield connection
        finally:
            connection.close()

    def event(self, kind, detail):
        with self.db() as db: db.execute('INSERT INTO events VALUES(?,?,?,?)',(identifier(),kind,detail,time.time()))

    def list_documents(self):
        with self.db() as db:
            return [dict(r) for r in db.execute('SELECT id,name,suffix,reviewed,pdf_id,created FROM documents ORDER BY created DESC,id DESC')]

    def document(self, document_id):
        with self.db() as db:
            row=db.execute('SELECT * FROM documents WHERE id=?',(document_id,)).fetchone()
            if not row: raise UserError('This document was not found.')
            value=dict(row)
            for key in ('meta','issues'): value[key]=json.loads(value[key])
            value['decisions']={r['issue_id']:r['note'] for r in db.execute('SELECT * FROM decisions WHERE document_id=?',(document_id,))}
            return value

    def path(self, document_id):
        doc=self.document(document_id)
        return self.blobs / (doc['id']+doc['suffix'])

    def import_document(self,name,data):
        name=safe_name(name)
        suffix=Path(name).suffix.lower()
        if suffix not in EXTENSIONS: raise UserError('Choose a .docx, .pdf or .txt file.')
        if not data or len(data)>MAX_FILE: raise UserError('Choose a non-empty file smaller than 20 MB.')
        sha=digest(data)
        with self.db() as db:
            existing=db.execute('SELECT id FROM documents WHERE hash=?',(sha,)).fetchone()
            if existing: return self.document(existing['id'])
        content,meta=extract(data,suffix)
        if len(content)>1_000_000: raise UserError('This document contains too much text. Split it into smaller documents.')
        doc_id=identifier()
        path=self.blobs/(doc_id+suffix)
        path.write_bytes(data)
        try:
            with self.db() as db:
                db.execute('INSERT INTO documents(id,name,suffix,hash,content,meta,issues,created) VALUES(?,?,?,?,?,?,?,?)',(doc_id,name,suffix,sha,content,json.dumps(meta),json.dumps(findings(content,meta)),time.time()))
        except sqlite3.IntegrityError:
            path.unlink()
            with self.db() as db: doc_id=db.execute('SELECT id FROM documents WHERE hash=?',(sha,)).fetchone()['id']
        self.event('document','Imported '+name)
        return self.document(doc_id)

    def decide(self,document_id,issue_id,note):
        doc=self.document(document_id)
        if issue_id not in {i['id'] for i in doc['issues']}: raise UserError('That review item no longer exists.')
        note=text(note,'review note',1000)
        with self.db() as db:
            db.execute('INSERT OR REPLACE INTO decisions VALUES(?,?,?)',(document_id,issue_id,note))
            db.execute('UPDATE documents SET reviewed=0 WHERE id=?',(document_id,))
        return self.document(document_id)

    def approve(self,document_id):
        doc=self.document(document_id)
        if any(i['severity']=='block' for i in doc['issues']): raise UserError('Resolve the blocking items in the original application, then import the revised document.')
        if any(i['id'] not in doc['decisions'] for i in doc['issues']): raise UserError('Review every flagged item before approving this document.')
        with self.db() as db: db.execute('UPDATE documents SET reviewed=1 WHERE id=?',(document_id,))
        self.event('review','Approved '+doc['name'])
        return self.document(document_id)

    def replace(self,document_id,old,new):
        doc=self.document(document_id)
        if doc['suffix'] not in {'.docx','.txt'}: raise UserError('Edit PDFs in the original document, then import the revised file.')
        old=text(old,'original text',500)
        new=text(new,'replacement',500,False)
        if old==new: raise UserError('The replacement must be different.')
        source=self.path(document_id).read_bytes()
        if doc['suffix']=='.txt':
            if doc['content'].count(old)!=1: raise UserError('Choose text that occurs exactly once.')
            output=doc['content'].replace(old,new,1).encode('utf-8')
        else:
            validate_docx(source)
            package=Document(io.BytesIO(source))
            paragraphs=list(package.paragraphs)
            def cells(tables):
                for table in tables:
                    for row in table.rows:
                        for cell in row.cells:
                            yield from cell.paragraphs
                            yield from cells(cell.tables)
            paragraphs.extend(cells(package.tables))
            matches=[p for p in paragraphs if old in p.text]
            if len(matches)!=1 or matches[0].text.count(old)!=1: raise UserError('Choose text that occurs once in the body or a table. Headers and footers should be edited in Word.')
            p=matches[0]
            runs=[r for r in p.runs if old in r.text]
            if len(runs)!=1: raise UserError('That text spans formatting changes. Edit it in Word to preserve the layout.')
            runs[0].text=runs[0].text.replace(old,new,1)
            stream=io.BytesIO();package.save(stream);output=stream.getvalue()
        name=Path(doc['name']).stem+' - revised'+doc['suffix']
        return self.import_document(name,output)

    def convert_to(self,document_id,target):
        doc=self.document(document_id)
        if target=='pdf':return self.convert(document_id,require_review=False)
        if target!='docx' or doc['suffix']!='.pdf':raise UserError('Choose Word to PDF or PDF to Word.')
        reader=PdfReader(io.BytesIO(self.path(document_id).read_bytes()))
        if len(reader.pages)>100:raise UserError('Choose a PDF with 100 pages or fewer.')
        paragraphs=[page.extract_text() or '' for page in reader.pages]
        if not any(p.strip() for p in paragraphs):raise UserError('This scanned PDF has no extractable text. OCR is needed before creating editable Word text.')
        if sum(map(len,paragraphs))>1000000:raise UserError('Choose a smaller PDF for conversion.')
        output=Document()
        for i,content in enumerate(paragraphs):
            if i:output.add_page_break()
            if not content.strip():output.add_paragraph('[No extractable text on this PDF page; check the original.]')
            for line in content.splitlines():output.add_paragraph(line)
        stream=io.BytesIO();output.save(stream)
        result=self.import_document(Path(doc['name']).stem+' - editable text.docx',stream.getvalue())
        self.event('convert','Created '+result['name']+'; text-only conversion requires review')
        return result

    def convert(self,document_id,require_review=True):
        doc=self.document(document_id)
        if require_review and not doc['reviewed']: raise UserError('Approve the document review first.')
        if doc['suffix']=='.pdf': return doc
        if doc['suffix']!='.docx': raise UserError('PDF conversion currently supports Word documents. Save text as .docx in Word first.')
        if doc['pdf_id']: return self.document(doc['pdf_id'])
        if not self.conversion_lock.acquire(blocking=False): raise UserError('Word is converting another document. Try again shortly.')
        target=self.blobs/(identifier()+'.pdf')
        try:
            validate_docx(self.path(document_id).read_bytes())
            script=Path(__file__).with_name('word-export.ps1')
            try:
                result=subprocess.run([shutil.which('pwsh') or 'powershell.exe','-NoProfile','-NonInteractive','-File',str(script),'-Source',str(self.path(document_id)),'-Destination',str(target)],capture_output=True,timeout=90,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
            except (OSError,subprocess.TimeoutExpired) as exc: raise UserError('Word conversion did not finish. Check that desktop Word is installed and activated, then retry.') from exc
            if result.returncode or not target.exists(): raise UserError('Word could not convert this file. Open it in Word to check activation, protection or repair prompts.')
            pdf=self.import_document(Path(doc['name']).stem+'.pdf',target.read_bytes())
            if not pdf['meta'].get('pages') or not pdf['content'].strip(): raise UserError('The PDF has no readable content. Inspect the original document in Word.')
            source_words=set(re.findall(r'\w+',doc['content'].casefold()))
            output_words=set(re.findall(r'\w+',pdf['content'].casefold()))
            if source_words and len(source_words-output_words)/len(source_words)>.15: raise UserError('The PDF text differs noticeably from the Word document. Inspect both files before sharing.')
            with self.db() as db: db.execute('UPDATE documents SET pdf_id=? WHERE id=?',(pdf['id'],document_id))
            self.event('pdf','Created '+pdf['name']+'; visual approval still required')
            return pdf
        finally:
            target.unlink(missing_ok=True)
            self.conversion_lock.release()

    def profile(self):
        with self.db() as db:
            row=db.execute("SELECT value FROM settings WHERE key='profile'").fetchone()
        return json.loads(row['value']) if row else None

    def save_profile(self,name,role='',preferences=''):
        profile={'name':text(name,'name',100),'role':text(role,'role',150,False),'preferences':text(preferences,'preferences',2000,False)}
        with self.db() as db:
            db.execute("INSERT OR REPLACE INTO settings(key,value) VALUES('profile',?)",(json.dumps(profile),))
        return profile

    def organise_file(self,file_id,folder_id,project):
        project=text(project,'project name',100)
        if project in ('.','..') or any(c in project for c in '<>:"/\\|?*') or project.endswith(('.', ' ')) or any(ord(c)<32 for c in project):
            raise UserError('Use a simple project folder name without slashes or special characters.')
        if project.split('.')[0].upper() in {'CON','PRN','AUX','NUL',*[f'COM{i}' for i in range(1,10)],*[f'LPT{i}' for i in range(1,10)]}:
            raise UserError('Choose another project folder name.')
        with self.db() as db:
            source=db.execute('SELECT files.*,folders.path AS root FROM files JOIN folders ON folder_id=folders.id WHERE files.id=?',(file_id,)).fetchone()
            folder=db.execute('SELECT path FROM folders WHERE id=?',(folder_id,)).fetchone()
        if not source or not folder: raise UserError('Refresh your folders and choose the file again.')
        original=Path(source['path']);root=Path(folder['path'])
        if original.is_symlink() or not original.resolve().is_relative_to(Path(source['root']).resolve()) or not original.is_file():
            raise UserError('The original file is no longer available in its work folder.')
        if not root.is_dir() or root.is_symlink() or getattr(root,'is_junction',lambda:False)(): raise UserError('The destination work folder is unavailable.')
        destination=root/project
        if destination.is_symlink() or getattr(destination,'is_junction',lambda:False)() or not destination.resolve().is_relative_to(root.resolve()):
            raise UserError('Choose a project folder inside the selected work folder.')
        data=original.read_bytes()
        if len(data)>MAX_FILE: raise UserError('Choose a file smaller than 20 MB.')
        destination.mkdir(exist_ok=True)
        target=destination/original.name
        try:
            with target.open('xb') as output: output.write(data)
        except FileExistsError: raise UserError('A file with this name already exists in that project. Nothing was overwritten.')
        self.event('organise','Copied '+original.name+' into project '+project)
        self.scan()
        return {'path':str(target),'original_kept':True}

    def add_folder(self,path,label):
        root=Path(text(path,'folder path',1000)).expanduser()
        if not root.is_absolute() or not root.is_dir(): raise UserError('Choose an existing absolute folder path on this computer.')
        root=root.resolve()
        if root==Path(root.anchor) or root==Path.home(): raise UserError('Choose a specific work folder, rather than an entire drive or user profile.')
        if root.is_relative_to(self.directory) or self.directory.is_relative_to(root): raise UserError('Choose your work documents folder, outside HARD’s application folder.')
        with self.db() as db:
            for row in db.execute('SELECT path FROM folders'):
                other=Path(row['path'])
                if root==other or root.is_relative_to(other) or other.is_relative_to(root): raise UserError('That folder overlaps a folder already included.')
            db.execute('INSERT INTO folders VALUES(?,?,?)',(identifier(),str(root),text(label,'folder label',100)))
        return self.scan()

    def scan(self):
        with self.db() as db: roots=[dict(r) for r in db.execute('SELECT * FROM folders')]
        records=[]; skipped=0;limited=False
        for root in roots:
            base=Path(root['path'])
            if not base.is_dir(): skipped+=1;continue
            for folder,dirs,names in os.walk(base,followlinks=False):
                dirs[:]=[n for n in dirs if not n.startswith('.') and not (Path(folder)/n).is_symlink() and not getattr(Path(folder)/n,'is_junction',lambda:False)()]
                for name in names:
                    path=Path(folder)/name
                    if path.suffix.lower() not in EXTENSIONS or name.startswith('~$'): continue
                    try:
                        real=path.resolve()
                        if path.is_symlink() or not real.is_relative_to(base): continue
                        stat=path.stat()
                        if stat.st_size>MAX_FILE: skipped+=1;continue
                        records.append((digest(str(real).casefold().encode())[:32],root['id'],str(real),name,stat.st_mtime,stat.st_size))
                    except OSError: skipped+=1
                    if len(records)>=10000: limited=True;break
                if limited: break
            if limited: break
        with self.db() as db:
            db.execute('DELETE FROM files')
            db.executemany('INSERT OR IGNORE INTO files VALUES(?,?,?,?,?,?)',records)
        return {'indexed':len(records),'skipped':skipped,'limited':limited}

    def search(self,query=''):
        query=text(query,'search',200,False).casefold()
        with self.db() as db:
            records=[dict(r) for r in db.execute('SELECT files.*,folders.label FROM files JOIN folders ON folder_id=folders.id ORDER BY modified DESC,files.id DESC')]
        terms=query.split()
        return [r for r in records if all(t in (r['name']+' '+r['path']).casefold() for t in terms)]

    def indexed_path(self,file_id):
        with self.db() as db:
            row=db.execute('SELECT files.*,folders.path AS root FROM files JOIN folders ON folder_id=folders.id WHERE files.id=?',(file_id,)).fetchone()
        if not row: raise UserError('Refresh your folders and choose the file again.')
        path=Path(row['path'])
        if path.is_symlink() or not path.resolve().is_relative_to(Path(row['root']).resolve()) or not path.is_file():
            raise UserError('The file is no longer available inside its work folder.')
        return path

    def rename_file(self,file_id,name,confirmed):
        if confirmed is not True: raise UserError('Confirm the new file name before renaming.')
        path=self.indexed_path(file_id);name=text(name,'file name',200)
        if name in ('.','..') or any(c in name for c in '<>:"/\\|?*') or name.endswith(('.', ' ')) or any(ord(c)<32 for c in name):
            raise UserError('Use a file name without slashes or special characters.')
        if name.split('.')[0].upper() in {'CON','PRN','AUX','NUL',*[f'COM{i}' for i in range(1,10)],*[f'LPT{i}' for i in range(1,10)]}:
            raise UserError('Choose another file name.')
        if Path(name).suffix.lower()!=path.suffix.lower(): raise UserError('Keep the original file extension.')
        target=path.with_name(name)
        if target.exists(): raise UserError('A file with this name already exists. Nothing was renamed.')
        if os.name!='nt': raise UserError('Renaming is available in the Windows desktop app.')
        try: path.rename(target)
        except OSError: raise UserError('The file could not be renamed. Close it in other applications and try again.') from None
        self.scan();self.event('rename','Renamed '+path.name+' to '+name)
        return {'id':digest(str(target.resolve()).casefold().encode())[:32],'previous_name':path.name,'name':name}

    def duplicates(self):
        records=self.search();sizes={};groups={};skipped=0
        for row in records:sizes.setdefault(row['size'],[]).append(row)
        for candidates in sizes.values():
            if len(candidates)<2:continue
            for row in candidates:
                try:
                    path=self.indexed_path(row['id'])
                    with path.open('rb') as source:data=source.read(MAX_FILE+1)
                    if len(data)>MAX_FILE:skipped+=1;continue
                    groups.setdefault(digest(data),[]).append(row)
                except (OSError,UserError):skipped+=1
        return {'groups':[{'hash':key,'files':rows} for key,rows in groups.items() if len(rows)>1],'skipped':skipped}

    def import_indexed(self,file_id):
        with self.db() as db:
            row=db.execute('SELECT files.*,folders.path AS root FROM files JOIN folders ON folder_id=folders.id WHERE files.id=?',(file_id,)).fetchone()
        if not row: raise UserError('This file is no longer indexed. Refresh your folders.')
        path=Path(row['path']).resolve()
        if not path.is_relative_to(Path(row['root']).resolve()) or path.is_symlink(): raise UserError('The file is outside the selected folder.')
        try:
            if path.stat().st_size>MAX_FILE: raise UserError('Choose a file smaller than 20 MB.')
            return self.import_document(path.name,path.read_bytes())
        except OSError as exc: raise UserError('This file is unavailable. For OneDrive files, make it available on this device first.') from exc

    def contacts(self):
        with self.db() as db: return [dict(r) for r in db.execute('SELECT * FROM contacts ORDER BY name')]

    def add_contact(self,name,email,organization,source,confirm):
        name=text(name,'name',100);email=text(email,'email address',254)
        if confirm is not True: raise UserError('Confirm the address against the original client request.')
        if not re.fullmatch(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9](?:[A-Za-z0-9.-]*[A-Za-z0-9])?\.[A-Za-z]{2,}",email) or '..' in email: raise UserError('Enter a valid email address without spaces or extra punctuation.')
        domain=email.rsplit('@',1)[1]
        if any(not s or s.startswith('-') or s.endswith('-') for s in domain.split('.')): raise UserError('Check the domain in this email address.')
        source=text(source,'verification source',500)
        contact={'id':identifier(),'name':name,'email':email,'organization':text(organization,'organisation',150,False),'source':source,'verified':time.time()}
        try:
            with self.db() as db: db.execute('INSERT INTO contacts VALUES(:id,:name,:email,:organization,:source,:verified)',contact)
        except sqlite3.IntegrityError as exc: raise UserError('That exact address is already saved.') from exc
        return contact

    def check_recipient(self,address):
        import difflib
        address=text(address,'recipient address',254)
        contacts=self.contacts()
        exact=next((c for c in contacts if c['email']==address),None)
        return {'verified':bool(exact),'contact':exact,'suggestions':[] if exact else [c for c in contacts if difflib.SequenceMatcher(None,c['email'].casefold(),address.casefold()).ratio()>.72][:5]}

    def stage_email(self,contact_id,document_id,subject,body,provider,visual_confirmed,recipient=None):
        contact=next((c for c in self.contacts() if c['id']==contact_id),None)
        if not contact and recipient is not None:
            if not isinstance(recipient,dict):raise UserError('Enter the recipient details.')
            address=text(recipient.get('email'),'email address',254)
            if recipient.get('confirmed') is not True:raise UserError('Check and confirm the typed recipient address.')
            if not re.fullmatch(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9](?:[A-Za-z0-9.-]*[A-Za-z0-9])?\.[A-Za-z]{2,}",address) or '..' in address or any(not part or part.startswith('-') or part.endswith('-') for part in address.rsplit('@',1)[-1].split('.')):
                raise UserError('Enter a valid email address without spaces or extra punctuation.')
            contact={'id':None,'name':text(recipient.get('name'),'recipient name',100),'email':address,'organization':'','verified':None,'source':'Typed and confirmed for this message'}
        if not contact: raise UserError('Choose a contact or enter a recipient.')
        doc=self.document(document_id)
        if doc['suffix']!='.pdf' or not doc['reviewed']: raise UserError('Review and approve the actual PDF attachment first.')
        if visual_confirmed is not True: raise UserError('Open the PDF and confirm its pages before preparing the email.')
        if provider not in {'gmail','outlook'}: raise UserError('Choose Gmail or Outlook.')
        subject=text(subject,'subject',200);body=text(body,'message',20000)
        if '\n' in subject or '\r' in subject: raise UserError('The subject must be one line.')
        payload={'contact':contact,'document_id':document_id,'attachment_name':doc['name'],'attachment_hash':doc['hash'],'subject':subject,'body':body,'provider':provider}
        draft_id=identifier()
        with self.db() as db: db.execute('INSERT INTO drafts(id,payload,status,created) VALUES(?,?,?,?)',(draft_id,json.dumps(payload),'draft',time.time()))
        return self.draft(draft_id)

    def draft(self,draft_id):
        with self.db() as db: row=db.execute('SELECT * FROM drafts WHERE id=?',(draft_id,)).fetchone()
        if not row: raise UserError('That draft was not found.')
        result=dict(row);result['payload']=json.loads(result['payload']);return result

    def export_eml(self,draft_id):
        from email.message import EmailMessage
        from email import policy
        draft=self.draft(draft_id);p=draft['payload']
        data=self.path(p['document_id']).read_bytes()
        if digest(data)!=p['attachment_hash']: raise UserError('The attachment changed. Prepare a new email.')
        msg=EmailMessage(policy=policy.SMTP)
        msg['To']=p['contact']['email'];msg['Subject']=p['subject'];msg['X-Unsent']='1'
        msg.set_content(p['body']);msg.add_attachment(data,maintype='application',subtype='pdf',filename=p['attachment_name'])
        return msg.as_bytes()

    def snapshot(self):
        with self.db() as db:
            return {'documents':self.list_documents(),'contacts':self.contacts(),'folders':[dict(r) for r in db.execute('SELECT * FROM folders')],'drafts':[dict(r) for r in db.execute('SELECT id,status,created FROM drafts ORDER BY created DESC,id DESC')],'events':[dict(r) for r in db.execute('SELECT * FROM events ORDER BY created DESC,id DESC')]}
