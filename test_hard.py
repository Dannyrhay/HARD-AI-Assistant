import base64
import io
import json
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch
from urllib.request import Request,urlopen
from urllib.error import HTTPError
from http.server import ThreadingHTTPServer
from docx import Document
from pypdf import PdfWriter
from hard_core import Workspace,UserError,safe_name
from hard_mail import Mail,protect
from server import handler_factory

class CoreTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.base=Path(self.tmp.name);self.w=Workspace(self.base/'data')
    def tearDown(self): self.tmp.cleanup()
    def word(self,value='Water assessment proposal'):
        d=Document();d.add_heading('Water assessment proposal',0);d.add_paragraph(value);b=io.BytesIO();d.save(b);return b.getvalue()
    def pdf(self):
        b=io.BytesIO();p=PdfWriter();p.add_blank_page(width=595,height=842);p.write(b);doc=self.w.import_document('Proposal.pdf',b.getvalue())
        # Blank fixture is explicitly approved only in isolated unit-test storage.
        with self.w.db() as db: db.execute("UPDATE documents SET reviewed=1 WHERE id=?",(doc['id'],))
        return self.w.document(doc['id'])
    def draft(self):
        doc=self.pdf();c=self.w.add_contact('Test Client','client@example.org','Test','Original request',True)
        return self.w.stage_email(c['id'],doc['id'],'Proposal','Please see the attachment.','gmail',True)
    def test_import_and_duplicate(self):
        data=self.word();a=self.w.import_document('Proposal.docx',data);b=self.w.import_document('Duplicate.docx',data);self.assertEqual(a['id'],b['id']);self.assertEqual(self.w.path(a['id']).read_bytes(),data)
    def test_placeholder_blocks_approval_until_reviewed(self):
        d=self.w.import_document('Proposal.docx',self.word('[INSERT DATE]'))
        with self.assertRaises(UserError): self.w.approve(d['id'])
        d=self.w.decide(d['id'],d['issues'][0]['id'],'Checked and intentionally retained for internal draft');self.assertTrue(self.w.approve(d['id'])['reviewed'])
    def test_revision_preserves_original(self):
        d=self.w.import_document('Proposal.docx',self.word('Client Alpha'));original=self.w.path(d['id']).read_bytes();r=self.w.replace(d['id'],'Client Alpha','Client Beta');self.assertIn('Client Beta',r['content']);self.assertEqual(self.w.path(d['id']).read_bytes(),original);self.assertFalse(r['reviewed'])
    def test_invalid_docx(self):
        with self.assertRaises(UserError): self.w.import_document('Bad.docx',b'bad zip')
    def test_path_names(self):
        for name in ['../file.txt','CON.txt','a/b.pdf','a\\b.pdf','wrong:thing.docx']:
            with self.assertRaises(UserError): safe_name(name)
    def test_missing_document(self):
        with self.assertRaises(UserError): self.w.document('../../anything')
    def test_block_scanned_pdf(self):
        d=self.pdf()
        with self.assertRaises(UserError): self.w.approve(d['id'])
    def test_folder_scope_and_search(self):
        folder=self.base/'files';folder.mkdir();(folder/'Kenya proposal.txt').write_text('Water proposal');self.w.add_folder(str(folder),'Work');rows=self.w.search('Kenya');self.assertEqual(len(rows),1);self.assertEqual(self.w.import_indexed(rows[0]['id'])['content'],'Water proposal')
        with self.assertRaises(UserError):self.w.add_folder(str(folder),'Duplicate')
    def test_no_whole_drive(self):
        with self.assertRaises(UserError):self.w.add_folder(str(self.base.anchor),'Drive')
    def test_contact_requires_confirmation(self):
        with self.assertRaises(UserError):self.w.add_contact('A','a@example.org','Firm','Request',False)
    def test_recipient_typo_and_case(self):
        self.w.add_contact('Mary','mary.otieno@example.org','Firm','Request',True);self.assertTrue(self.w.check_recipient('mary.otieno@example.org')['verified']);r=self.w.check_recipient('mary.ootieno@example.org');self.assertFalse(r['verified']);self.assertEqual(len(r['suggestions']),1);self.assertFalse(self.w.check_recipient('MARY.otieno@example.org')['verified'])
    def test_header_injection(self):
        with self.assertRaises(UserError):self.w.add_contact('Mary','a@example.org\r\nBcc: other@example.org','Firm','Request',True)
    def test_draft_attachment_and_persistence(self):
        d=self.draft();eml=self.w.export_eml(d['id']);self.assertIn(b'client@example.org',eml);self.assertIn(b'application/pdf',eml);self.assertEqual(Workspace(self.base/'data').draft(d['id'])['payload'],d['payload'])
    def test_attachment_change_invalidates_export(self):
        d=self.draft();self.w.path(d['payload']['document_id']).write_bytes(b'changed')
        with self.assertRaises(UserError):self.w.export_eml(d['id'])
    def test_unreviewed_attachment_rejected(self):
        d=self.draft();p=d['payload']
        with self.w.db() as db:db.execute('UPDATE documents SET reviewed=0 WHERE id=?',(p['document_id'],))
        with self.assertRaises(UserError):self.w.stage_email(p['contact']['id'],p['document_id'],'Subject','Body','gmail',True)
    def test_send_single_attempt_and_replay(self):
        d=self.draft();mail=Mail(self.w,'http://127.0.0.1:5188');token={'access_token':'fake','address':'sender@example.org'}
        with patch.object(mail,'access',return_value=token),patch('hard_mail.request_json',return_value={'id':'receipt'}) as send:
            a=mail.approve(d['id'],True,'sender@example.org');self.assertEqual(mail.send(d['id'],a['approval'])['status'],'sent')
            with self.assertRaises(UserError):mail.send(d['id'],a['approval'])
            self.assertEqual(send.call_count,1)
    def test_timeout_is_uncertain_not_retryable(self):
        d=self.draft();mail=Mail(self.w,'http://127.0.0.1:5188')
        with patch.object(mail,'access',return_value={'access_token':'fake','address':'sender@example.org'}),patch('hard_mail.request_json',side_effect=TimeoutError):
            a=mail.approve(d['id'],True,'sender@example.org')
            with self.assertRaises(UserError):mail.send(d['id'],a['approval'])
            self.assertEqual(self.w.draft(d['id'])['status'],'uncertain')
            with self.assertRaises(UserError):mail.approve(d['id'],True,'sender@example.org')
    def test_changed_sender_rejected(self):
        d=self.draft();mail=Mail(self.w,'http://127.0.0.1:5188')
        with patch.object(mail,'access',return_value={'access_token':'fake','address':'different@example.org'}):
            with self.assertRaises(UserError):mail.approve(d['id'],True,'sender@example.org')
    def test_expired_approval_rejected(self):
        d=self.draft();mail=Mail(self.w,'http://127.0.0.1:5188')
        with patch.object(mail,'access',return_value={'access_token':'fake','address':'sender@example.org'}):
            a=mail.approve(d['id'],True,'sender@example.org');mail.approvals[a['approval']]['expires']=0
            with self.assertRaises(UserError):mail.send(d['id'],a['approval'])
    def test_tracked_changes_block(self):
        import zipfile
        raw=self.word('Water proposal');target=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as src,zipfile.ZipFile(target,'w') as out:
            for name in src.namelist():
                data=src.read(name)
                if name=='word/document.xml':data=data.replace(b'<w:body>',b'<w:body><w:ins><w:r><w:t>Inserted text</w:t></w:r></w:ins>')
                out.writestr(name,data)
        d=self.w.import_document('Tracked.docx',target.getvalue())
        with self.assertRaises(UserError):self.w.approve(d['id'])
    def test_corrupt_pdf_rejected(self):
        with self.assertRaises(UserError):self.w.import_document('Bad.pdf',b'not PDF')
    def test_oauth_pkce_and_single_use_state(self):
        from urllib.parse import urlparse,parse_qs
        mail=Mail(self.w,'http://127.0.0.1:5188')
        with patch.object(mail,'config',return_value={'client_id':'test-client','client_secret':''}):
            result=mail.begin('gmail');query=parse_qs(urlparse(result['url']).query);self.assertEqual(query['code_challenge_method'],['S256']);self.assertEqual(query['redirect_uri'],['http://127.0.0.1:5188/oauth/callback'])
            with patch('hard_mail.request_json',side_effect=[{'access_token':'fake','refresh_token':'fake','expires_in':3600},{'emailAddress':'sender@example.org'}]):mail.callback(query['state'][0],'fake-code')
            with self.assertRaises(UserError):mail.callback(query['state'][0],'fake-code')
    def test_google_setup_encrypted_and_not_in_status(self):
        mail=Mail(self.w,'http://127.0.0.1:5188')
        credentials={'installed':{'client_id':'123-test.apps.googleusercontent.com','client_secret':'test-client-secret','auth_uri':'https://untrusted.example/'}}
        with patch.dict('os.environ',{},clear=True):
            self.assertTrue(mail.configure('gmail',credentials)['configured'])
            self.assertEqual(mail.config('gmail')['client_secret'],'test-client-secret')
            self.assertNotIn(b'test-client-secret',(self.w.directory/'gmail.client').read_bytes())
            status=mail.status();self.assertTrue(status['gmail']['configured']);self.assertFalse(status['gmail']['connected']);self.assertNotIn('test-client-secret',json.dumps(status))
            self.assertTrue(mail.begin('gmail')['url'].startswith('https://accounts.google.com/'))
    def test_google_web_client_rejected(self):
        mail=Mail(self.w,'http://127.0.0.1:5188')
        with self.assertRaises(UserError):mail.configure('gmail',{'web':{'client_id':'123-test.apps.googleusercontent.com'}})
    def test_setup_connected_account_not_replaced(self):
        mail=Mail(self.w,'http://127.0.0.1:5188');(self.w.directory/'gmail.token').write_bytes(b'fixture')
        with patch.dict('os.environ',{},clear=True):
            with self.assertRaises(UserError):mail.configure('gmail',{'installed':{'client_id':'123-test.apps.googleusercontent.com'}})
    def test_microsoft_setup_public_client(self):
        mail=Mail(self.w,'http://127.0.0.1:5188')
        with patch.dict('os.environ',{},clear=True):
            mail.configure('outlook',{'client_id':'00000000-0000-4000-8000-000000000000','client_secret':'must-not-store'})
            self.assertEqual(mail.config('outlook')['client_secret'],'')
            with self.assertRaises(UserError):mail.configure('outlook',{'client_id':'not-a-uuid'})
    def test_environment_config_not_silently_replaced(self):
        mail=Mail(self.w,'http://127.0.0.1:5188')
        with patch.dict('os.environ',{'HARD_GOOGLE_CLIENT_ID':'existing'}):
            with self.assertRaises(UserError):mail.configure('gmail',{'installed':{'client_id':'123-test.apps.googleusercontent.com'}})
    def test_oauth_invalid_state(self):
        with self.assertRaises(UserError):Mail(self.w,'http://127.0.0.1:5188').callback('bad','bad')
    def test_secret_encryption(self):
        raw=b'fake-test-token';encrypted=protect(raw);self.assertNotIn(raw,encrypted);self.assertEqual(protect(encrypted,True),raw)

class HTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory();cls.w=Workspace(Path(cls.tmp.name)/'data');cls.server=ThreadingHTTPServer(('127.0.0.1',0),lambda *a:None);cls.origin='http://127.0.0.1:'+str(cls.server.server_port);cls.mail=Mail(cls.w,cls.origin);cls.server.RequestHandlerClass=handler_factory(cls.w,cls.mail,cls.origin);cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True);cls.thread.start()
    @classmethod
    def tearDownClass(cls):cls.server.shutdown();cls.server.server_close();cls.thread.join();cls.tmp.cleanup()
    def test_state_and_static(self):
        with urlopen(self.origin+'/api/state') as r:self.assertTrue(json.load(r)['csrf'])
        with urlopen(self.origin+'/') as r:self.assertIn(b'HARD',r.read());self.assertIn("frame-ancestors 'none'",r.headers['Content-Security-Policy'])
    def test_csrf_rejects(self):
        req=Request(self.origin+'/api/scan',data=b'{}',headers={'Content-Type':'application/json','Origin':'https://evil.example'})
        with self.assertRaises(HTTPError) as result:urlopen(req)
        self.assertEqual(result.exception.code,400)
    def test_host_rejects(self):
        with self.assertRaises(HTTPError):urlopen(Request(self.origin+'/api/state',headers={'Host':'evil.example'}))
    def test_data_not_served(self):
        with self.assertRaises(HTTPError) as result:urlopen(self.origin+'/.data/workspace.sqlite3')
        self.assertEqual(result.exception.code,404)
    def test_real_http_import(self):
        with urlopen(self.origin+'/api/state') as r:csrf=json.load(r)['csrf']
        body=json.dumps({'name':'HTTP.txt','data':base64.b64encode(b'Water report').decode()}).encode();req=Request(self.origin+'/api/import',data=body,headers={'Content-Type':'application/json','Origin':self.origin,'X-HARD-CSRF':csrf})
        with urlopen(req) as r:self.assertEqual(json.load(r)['content'],'Water report')

if __name__=='__main__':unittest.main(verbosity=2)
